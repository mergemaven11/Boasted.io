"""Regression tests for the expanded career packet catalog and file exports."""
from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO

import mongomock
import pytest
from bson import ObjectId
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfReader

import app.interview_packet_routes as interview_routes
import app.packet_audit as packet_audit
import app.packet_audit_routes as packet_audit_routes
import app.packet_catalog_routes as catalog_routes
import app.packet_platform_routes as platform_routes
import app.performance_packet_routes as performance_routes
from app.main import app


client = TestClient(app, base_url="https://testserver")
PACKET_TYPES = [
    "performance-review",
    "promotion",
    "interview",
    "certification",
    "program-application",
    "scholarship",
    "portfolio",
    "career-transition",
]


@pytest.fixture
def packet_catalog_context(monkeypatch):
    mock_client = mongomock.MongoClient()
    db = mock_client["bragstack_packet_catalog_test"]
    entries = db["entries"]
    receipts = db["impact_receipts"]
    audits = db["packet_export_audit"]

    monkeypatch.setattr(performance_routes, "entries_collection", entries)
    monkeypatch.setattr(performance_routes, "impact_receipts_collection", receipts)
    monkeypatch.setattr(platform_routes, "impact_receipts_collection", receipts)
    monkeypatch.setattr(interview_routes, "impact_receipts_collection", receipts)
    monkeypatch.setattr(packet_audit, "packet_export_audit_collection", audits)
    monkeypatch.setattr(packet_audit_routes, "packet_export_audit_collection", audits)

    user = {
        "_id": ObjectId(),
        "name": "Jordan Lee",
        "email": "jordan@example.com",
        "headline": "Platform Support Engineer",
        "location": "Atlanta, Georgia",
        "plan": "pro",
    }
    app.dependency_overrides[catalog_routes.get_current_user] = lambda: user

    user_id = str(user["_id"])
    fixtures = [
        {
            "title": "Automated a recurring support validation workflow",
            "category": "Automation",
            "date": "2026-04-10",
            "result": "Reduced a repeated validation process from 45 minutes to 12 minutes.",
            "skills": ["Python", "Automation", "Troubleshooting"],
        },
        {
            "title": "Improved cross-team incident handoff documentation",
            "category": "Operations",
            "date": "2026-05-20",
            "result": "Created a reusable handoff pattern used across three support rotations.",
            "skills": ["Documentation", "Incident Response", "Communication"],
        },
    ]
    entry_ids = []
    for index, fixture in enumerate(fixtures):
        inserted = entries.insert_one(
            {
                "user_id": user_id,
                "title": fixture["title"],
                "category": fixture["category"],
                "entry_type": "Current Job",
                "entry_date": fixture["date"],
                "tags": fixture["skills"],
                "impact": fixture["result"],
                "resume_bullet": fixture["title"],
                "created_at": datetime.now(timezone.utc),
            }
        )
        entry_id = str(inserted.inserted_id)
        entry_ids.append(entry_id)
        receipts.insert_one(
            {
                "user_id": user_id,
                "source_entry_id": entry_id,
                "accomplishment": fixture["title"],
                "contribution": "Owned the documented work and follow-through.",
                "result": fixture["result"],
                "evidence": [
                    {
                        "title": f"Supporting artifact {index + 1}",
                        "evidence_type": "documentation",
                        "reference": f"EVIDENCE-{index + 1}",
                        "description": "Supporting proof for the saved accomplishment.",
                        "is_public": False,
                    }
                ],
                "skills": fixture["skills"],
                "credit": [],
                "confirmations": [
                    {
                        "name": "Avery Reviewer",
                        "role": "Peer",
                        "confirmation_type": "peer",
                        "status": "confirmed",
                    }
                ],
                "trust_signals": ["self-documented", "evidence-linked"],
                "is_public": False,
                "created_at": datetime.now(timezone.utc),
                "updated_at": datetime.now(timezone.utc),
            }
        )

    yield {"entries": entries, "receipts": receipts, "audits": audits, "user": user, "entry_ids": entry_ids}
    app.dependency_overrides.clear()


def _payload(packet_type: str, entry_ids: list[str]) -> dict:
    common = {
        "career_area": "Technology",
        "role_title": "Platform Support Engineer",
        "organization": "Example Technology Team",
        "theme": "modern-minimal",
        "brand_name": "Jordan Lee Career Proof",
        "reviewer_name": "Example Reviewer",
        "signature_entry_ids": entry_ids,
    }
    additions = {
        "performance-review": {"review_cycle_label": "2026 Annual Review"},
        "promotion": {"target_role": "Platform Engineer", "target_level": "Level II"},
        "interview": {"target_role": "Platform Engineer", "target_organization": "Example Cloud Company", "selected_entry_ids": entry_ids},
        "certification": {"credential_name": "Cloud Operations Certification", "issuing_body": "Example Certification Board", "review_type": "Certification Review", "requirement_notes": "Demonstrate applied operations and troubleshooting experience."},
        "program-application": {"program_name": "Cloud Engineering Fellowship", "institution_name": "Example Institute", "application_type": "Fellowship", "application_deadline": "2026-11-01", "application_prompt": "Describe applied technical growth and impact."},
        "scholarship": {"scholarship_name": "Technology Leadership Scholarship", "sponsor_name": "Example Foundation", "award_focus": "Leadership and technical growth", "scholarship_deadline": "2026-12-01", "essay_prompt": "Describe how your work created measurable impact."},
        "portfolio": {"portfolio_title": "Platform Engineering Portfolio", "portfolio_audience": "Hiring Manager", "portfolio_focus": "Automation and reliability", "project_notes": "Emphasize reusable automation and operational documentation."},
        "career-transition": {"target_industry": "Cloud Platform Engineering", "target_role": "Platform Engineer", "transition_goal": "Move from support engineering into platform engineering.", "transferable_skills_focus": "Troubleshooting, automation, incident response, documentation", "transition_notes": "Show how support work demonstrates platform engineering capabilities."},
    }
    return {**common, **additions[packet_type]}


def _pdf_text(data: bytes) -> str:
    reader = PdfReader(BytesIO(data))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _docx_text(data: bytes) -> str:
    document = Document(BytesIO(data))
    text = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            text.extend(cell.text for cell in row.cells)
    return "\n".join(text)


@pytest.mark.parametrize("packet_type", PACKET_TYPES)
def test_catalog_packet_generation_contains_saved_proof(packet_catalog_context, packet_type):
    response = client.post(
        f"/packets/catalog/{packet_type}",
        json=_payload(packet_type, packet_catalog_context["entry_ids"]),
    )
    assert response.status_code == 200, response.text
    packet = response.json()["packet"]
    assert packet["kind"] == packet_type
    assert packet["title"] == catalog_routes.PACKET_SPECS[packet_type]["title"]
    assert packet["subject"]["name"] == "Jordan Lee"
    assert packet["scorecard"]["accomplishments"] == 2
    assert packet["scorecard"]["impact_receipts"] == 2
    assert packet["quality_message"] == catalog_routes.QUALITY_MESSAGE
    assert packet["usage_example"]
    assert packet["signature_accomplishments"]
    assert len(packet["reviewer_guide"]["criteria"]) == 6
    assert packet["reviewer_guide"]["feedback_sections"] == [
        "Strengths / what stands out",
        "Corrections or factual changes",
        "Questions / clarification needed",
        "Recommended next steps",
    ]
    assert "do not automatically change" in packet["reviewer_guide"]["source_record_notice"]


@pytest.mark.parametrize("packet_type", PACKET_TYPES)
def test_catalog_pdf_download_is_nonblank_and_contains_packet_content(packet_catalog_context, packet_type):
    response = client.post(
        f"/packets/catalog/{packet_type}.pdf",
        json=_payload(packet_type, packet_catalog_context["entry_ids"]),
    )
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("application/pdf")
    assert response.content.startswith(b"%PDF")
    assert len(response.content) > 2500
    text = _pdf_text(response.content)
    assert "Jordan Lee" in text
    assert catalog_routes.PACKET_SPECS[packet_type]["title"] in text
    assert "Impact Receipts" in text
    assert "Automated a recurring support validation workflow" in text
    assert "REVIEWER WORKSHEET" in text
    assert "Corrections or factual changes" in text
    assert "not automatic edits" in text


@pytest.mark.parametrize("packet_type", PACKET_TYPES)
def test_catalog_docx_download_is_nonblank_and_contains_packet_content(packet_catalog_context, packet_type):
    response = client.post(
        f"/packets/catalog/{packet_type}.docx",
        json=_payload(packet_type, packet_catalog_context["entry_ids"]),
    )
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert response.content.startswith(b"PK")
    assert len(response.content) > 5000
    text = _docx_text(response.content)
    assert "Jordan Lee" in text
    assert catalog_routes.PACKET_SPECS[packet_type]["title"] in text
    assert "Impact Receipts" in text
    assert "Automated a recurring support validation workflow" in text
    assert "REVIEWER WORKSHEET" in text
    assert "Corrections or factual changes" in text
    assert "not automatic edits" in text
