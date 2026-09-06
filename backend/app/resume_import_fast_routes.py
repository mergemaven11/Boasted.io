"""Optimized resume import and structured resume persistence routes."""
from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, status
from pydantic import BaseModel, Field
from starlette.concurrency import run_in_threadpool

from app.auth import get_current_user
from app.database import resume_documents_collection
from app.plans import require_feature
from app.resume_builder_routes import (
    MAX_RESUME_BYTES,
    MAX_RESUME_PAGES,
    ResumeBulletPayload,
    ResumeContactPayload,
    ResumeExperiencePayload,
    _server_readiness,
    _validate_saved_bullet_sources,
)
from app.resume_import_parser_v2 import parse_existing_resume_text

router = APIRouter(prefix="/resume-builder", tags=["resume-builder"])


class ResumeSupportingSectionsV2(BaseModel):
    projects: list[str] = Field(default_factory=list, max_length=50)
    education: list[str] = Field(default_factory=list, max_length=50)
    certifications: list[str] = Field(default_factory=list, max_length=50)
    leadership: list[str] = Field(default_factory=list, max_length=50)
    volunteer: list[str] = Field(default_factory=list, max_length=50)
    awards: list[str] = Field(default_factory=list, max_length=50)
    publications: list[str] = Field(default_factory=list, max_length=50)
    languages: list[str] = Field(default_factory=list, max_length=50)


class ResumeSaveRequestV2(BaseModel):
    title: str = Field(min_length=2, max_length=160)
    target_role: str = Field(default="", max_length=120)
    job_description: str = Field(default="", max_length=20000)
    summary: str = Field(default="", max_length=1200)
    bullets: list[ResumeBulletPayload] = Field(default_factory=list, max_length=100)
    skills: list[str] = Field(default_factory=list, max_length=80)
    contact: ResumeContactPayload = Field(default_factory=ResumeContactPayload)
    experience: list[ResumeExperiencePayload] = Field(default_factory=list, max_length=20)
    supporting_sections: ResumeSupportingSectionsV2 = Field(default_factory=ResumeSupportingSectionsV2)
    template_id: str = Field(default="classic-navy", max_length=80)


def _pdf_text_fast(data: bytes) -> str:
    """Use PyMuPDF first and invoke the slower fallback only for sparse output."""
    primary = ""
    try:
        import pymupdf
        lines: list[str] = []
        with pymupdf.open(stream=data, filetype="pdf") as document:
            if document.page_count > MAX_RESUME_PAGES:
                raise HTTPException(status_code=400, detail=f"Resume PDFs can contain up to {MAX_RESUME_PAGES} pages")
            for page in document:
                for *_, text, _block_no, _block_type in page.get_text("blocks", sort=True):
                    cleaned = "\n".join(part.strip() for part in str(text or "").splitlines() if part.strip())
                    if cleaned:
                        lines.append(cleaned)
        primary = "\n".join(lines).strip()
    except HTTPException:
        raise
    except Exception:
        primary = ""

    if len(primary) >= 220 and len(primary.splitlines()) >= 6:
        return primary

    try:
        from pypdf import PdfReader
        reader = PdfReader(BytesIO(data))
        if len(reader.pages) > MAX_RESUME_PAGES:
            raise HTTPException(status_code=400, detail=f"Resume PDFs can contain up to {MAX_RESUME_PAGES} pages")
        secondary = "\n".join((page.extract_text() or "") for page in reader.pages).strip()
    except HTTPException:
        raise
    except Exception:
        secondary = ""

    return secondary if len(secondary) > len(primary) else primary


def _extract_fast(filename: str, content_type: str, data: bytes) -> str:
    suffix = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if content_type == "application/pdf" or suffix == ".pdf":
        return _pdf_text_fast(data)
    if content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document" or suffix == ".docx":
        from docx import Document
        document = Document(BytesIO(data))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        table_rows = [" | ".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows]
        return "\n".join([*paragraphs, *table_rows]).strip()
    if content_type == "text/plain" or suffix == ".txt":
        return data.decode("utf-8", errors="replace").strip()
    raise HTTPException(status_code=415, detail="Upload a PDF, DOCX, or TXT resume")


def _parse_file(filename: str, content_type: str, data: bytes) -> dict:
    text = _extract_fast(filename, content_type, data).strip()
    if len(text) < 20:
        raise HTTPException(
            status_code=400,
            detail="We could not find enough readable text in that resume. If it is a scanned/image-only PDF, export a text-searchable PDF or DOCX and try again.",
        )
    parsed = parse_existing_resume_text(text)
    return {
        "filename": filename,
        "text": parsed["text"],
        "raw_text": parsed.get("raw_text", text),
        "summary": parsed["summary"],
        "bullets": parsed["bullets"],
        "skills": parsed["skills"],
        "sections_found": parsed["sections_found"],
        "sections": parsed.get("sections", {}),
        "header_lines": parsed.get("header_lines", []),
        "contact": parsed.get("contact", {}),
        "experience": parsed.get("experience", []),
        "parse_warnings": parsed.get("parse_warnings", []),
        "source_signals": parsed.get("source_signals", {}),
        "field_confidence": parsed.get("field_confidence", {}),
        "parse_quality": parsed.get("parse_quality", {}),
        "line_count": parsed["line_count"],
    }


def _saved_resume_view(item: dict) -> dict:
    return {
        "id": str(item["_id"]),
        "title": item.get("title", "Untitled resume"),
        "target_role": item.get("target_role", ""),
        "job_description": item.get("job_description", ""),
        "summary": item.get("summary", ""),
        "bullets": item.get("bullets", []),
        "skills": item.get("skills", []),
        "contact": item.get("contact", {}),
        "experience": item.get("experience", []),
        "supporting_sections": item.get("supporting_sections", {}),
        "template_id": item.get("template_id", "classic-navy"),
        "readiness": item.get("readiness", {}),
        "schema_version": item.get("schema_version", 3),
        "created_at": item.get("created_at"),
        "updated_at": item.get("updated_at"),
    }


@router.post("/import-fast")
async def import_resume_fast(file: UploadFile = File(...), current_user: dict = Depends(get_current_user)):
    require_feature(current_user, "resume_builder")
    filename = (file.filename or "resume").strip()[:240]
    data = await file.read(MAX_RESUME_BYTES + 1)
    if len(data) > MAX_RESUME_BYTES:
        raise HTTPException(status_code=413, detail="Resume file must be 5 MB or smaller")
    if not data:
        raise HTTPException(status_code=400, detail="Resume file is empty")
    try:
        return await run_in_threadpool(_parse_file, filename, file.content_type or "", data)
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=400, detail="We could not read that resume. Try exporting it as a fresh PDF or DOCX.") from exc


@router.post("/resumes-v2", status_code=status.HTTP_201_CREATED)
def save_resume_v2(payload: ResumeSaveRequestV2, current_user: dict = Depends(get_current_user)):
    """Persist the structured master resume without dropping optional sections."""
    require_feature(current_user, "resume_builder")
    user_id = str(current_user["_id"])
    canonical_bullets = [bullet for role in payload.experience for bullet in role.bullets] if payload.experience else payload.bullets
    if len(canonical_bullets) > 100:
        raise HTTPException(status_code=400, detail="A saved resume can contain up to 100 experience bullets")
    _validate_saved_bullet_sources(user_id, canonical_bullets)

    now = datetime.now(timezone.utc)
    document = {
        "user_id": user_id,
        "title": payload.title.strip(),
        "target_role": payload.target_role.strip(),
        "job_description": payload.job_description,
        "summary": payload.summary,
        "bullets": [bullet.model_dump() for bullet in payload.bullets],
        "skills": [skill.strip()[:120] for skill in payload.skills if skill.strip()],
        "contact": payload.contact.model_dump(),
        "experience": [role.model_dump() for role in payload.experience],
        "supporting_sections": payload.supporting_sections.model_dump(),
        "template_id": payload.template_id.strip() or "classic-navy",
        "readiness": _server_readiness(canonical_bullets),
        "schema_version": 5,
        "created_at": now,
        "updated_at": now,
    }
    result = resume_documents_collection.insert_one(document)
    return {"id": str(result.inserted_id), **{key: value for key, value in document.items() if key != "user_id"}}


@router.get("/resumes-v2")
def list_resumes_v2(current_user: dict = Depends(get_current_user)):
    require_feature(current_user, "resume_builder")
    cursor = resume_documents_collection.find({"user_id": str(current_user["_id"])}).sort("updated_at", -1).limit(50)
    return {"resumes": [_saved_resume_view(item) for item in cursor]}
