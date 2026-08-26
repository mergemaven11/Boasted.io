from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO

import fitz
from bson import ObjectId
from docx import Document
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, status
from pydantic import BaseModel, Field
from pypdf import PdfReader

from app.auth import get_current_user
from app.database import impact_receipts_collection, resume_documents_collection
from app.plans import require_feature
from app.resume_builder import analyze_resume
from app.resume_import_parser import parse_existing_resume_text

router = APIRouter(prefix="/resume-builder", tags=["resume-builder"])
MAX_RESUME_BYTES = 5 * 1024 * 1024
MAX_RESUME_PAGES = 12


class ResumeBuildRequest(BaseModel):
    target_role: str = Field(min_length=2, max_length=120)
    job_description: str = Field(min_length=20, max_length=20000)
    selected_receipt_ids: list[str] = Field(default_factory=list, max_length=100)
    existing_resume_text: str = Field(default="", max_length=50000)


class ResumeBulletPayload(BaseModel):
    text: str = Field(min_length=1, max_length=500)
    source_receipt_id: str = Field(default="", max_length=64)
    source_title: str = Field(default="Impact Receipt", max_length=240)
    source_kind: str = Field(default="impact-receipt", pattern="^(impact-receipt|imported|manual)$")
    matched_terms: list[str] = Field(default_factory=list, max_length=50)
    evidence_count: int = Field(default=0, ge=0, le=1000)
    has_metrics: bool = False
    edited: bool = False


class ResumeContactPayload(BaseModel):
    name: str = Field(default="", max_length=160)
    email: str = Field(default="", max_length=254)
    phone: str = Field(default="", max_length=80)
    location: str = Field(default="", max_length=160)
    linkedin: str = Field(default="", max_length=500)
    github: str = Field(default="", max_length=500)


class ResumeExperiencePayload(BaseModel):
    company: str = Field(min_length=1, max_length=240)
    title: str = Field(min_length=1, max_length=240)
    location: str = Field(default="", max_length=160)
    start_date: str = Field(default="", max_length=80)
    end_date: str = Field(default="", max_length=80)
    current: bool = False
    dates_raw: str = Field(default="", max_length=180)
    confidence: str = Field(default="low", pattern="^(high|medium|low|manual)$")
    bullets: list[ResumeBulletPayload] = Field(default_factory=list, max_length=20)


class ResumeSupportingSectionsPayload(BaseModel):
    projects: list[str] = Field(default_factory=list, max_length=50)
    education: list[str] = Field(default_factory=list, max_length=50)


class ResumeSaveRequest(BaseModel):
    title: str = Field(min_length=2, max_length=160)
    target_role: str = Field(default="", max_length=120)
    job_description: str = Field(default="", max_length=20000)
    summary: str = Field(default="", max_length=1200)
    bullets: list[ResumeBulletPayload] = Field(default_factory=list, max_length=20)
    skills: list[str] = Field(default_factory=list, max_length=40)
    contact: ResumeContactPayload = Field(default_factory=ResumeContactPayload)
    experience: list[ResumeExperiencePayload] = Field(default_factory=list, max_length=20)
    supporting_sections: ResumeSupportingSectionsPayload = Field(default_factory=ResumeSupportingSectionsPayload)


def _owned_receipts(user_id: str, selected_ids: list[str]) -> list[dict]:
    query: dict = {"user_id": user_id}
    if selected_ids:
        valid_ids = [ObjectId(value) for value in selected_ids if ObjectId.is_valid(value)]
        if not valid_ids:
            return []
        query["_id"] = {"$in": valid_ids}
    return list(impact_receipts_collection.find(query).sort("created_at", -1).limit(100))


def _validate_saved_bullet_sources(user_id: str, bullets: list[ResumeBulletPayload]) -> None:
    receipt_bullets = [bullet for bullet in bullets if bullet.source_kind == "impact-receipt"]
    raw_ids = {bullet.source_receipt_id for bullet in receipt_bullets}
    source_ids = {value for value in raw_ids if ObjectId.is_valid(value)}
    if len(source_ids) != len(raw_ids):
        raise HTTPException(status_code=400, detail="Every source-linked bullet must reference a valid Impact Receipt")
    if not source_ids:
        return
    owned_count = impact_receipts_collection.count_documents({"_id": {"$in": [ObjectId(value) for value in source_ids]}, "user_id": user_id})
    if owned_count != len(source_ids):
        raise HTTPException(status_code=400, detail="Resume bullet source does not belong to this account")


def _canonical_saved_bullets(payload: ResumeSaveRequest) -> list[ResumeBulletPayload]:
    structured = [bullet for role in payload.experience for bullet in role.bullets]
    return structured if payload.experience else payload.bullets


def _server_readiness(bullets: list[ResumeBulletPayload]) -> dict:
    edited_count = sum(1 for bullet in bullets if bullet.edited)
    imported_count = sum(1 for bullet in bullets if bullet.source_kind == "imported")
    source_linked_count = sum(1 for bullet in bullets if bullet.source_kind == "impact-receipt")
    return {
        "source_linked_draft": edited_count == 0 and imported_count == 0,
        "edited_bullets_needing_review": edited_count,
        "imported_bullets": imported_count,
        "source_linked_bullets": source_linked_count,
        "verification_status": "source-linked" if edited_count == 0 and imported_count == 0 else "mixed-sources",
    }


def _extract_pdf_text_pymupdf(data: bytes) -> str:
    """Primary born-digital PDF extraction with block geometry and stable reading-order sorting."""
    lines: list[str] = []
    with fitz.open(stream=data, filetype="pdf") as document:
        if document.page_count > MAX_RESUME_PAGES:
            raise HTTPException(status_code=400, detail=f"Resume PDFs can contain up to {MAX_RESUME_PAGES} pages")
        for page in document:
            blocks = page.get_text("blocks", sort=True)
            page_lines: list[str] = []
            for x0, y0, x1, y1, text, *_ in blocks:
                del x0, y0, x1, y1
                cleaned = "\n".join(part.strip() for part in str(text or "").splitlines() if part.strip())
                if cleaned:
                    page_lines.append(cleaned)
            lines.extend(page_lines)
    return "\n".join(lines).strip()


def _extract_pdf_text_pypdf(data: bytes) -> str:
    reader = PdfReader(BytesIO(data))
    if len(reader.pages) > MAX_RESUME_PAGES:
        raise HTTPException(status_code=400, detail=f"Resume PDFs can contain up to {MAX_RESUME_PAGES} pages")
    return "\n".join((page.extract_text() or "") for page in reader.pages).strip()


def _extract_pdf_text(data: bytes) -> str:
    """Use two independent extractors and keep the richer result without inventing content."""
    primary = ""
    secondary = ""
    try:
        primary = _extract_pdf_text_pymupdf(data)
    except HTTPException:
        raise
    except Exception:
        primary = ""

    try:
        secondary = _extract_pdf_text_pypdf(data)
    except HTTPException:
        raise
    except Exception:
        secondary = ""

    # Prefer geometry-aware PyMuPDF when it recovered comparable content. If one
    # extractor clearly lost text, use the richer source rather than silently
    # accepting the sparse interpretation.
    if primary and (not secondary or len(primary) >= len(secondary) * 0.72):
        return primary
    return secondary or primary


def _extract_uploaded_text(filename: str, content_type: str, data: bytes) -> str:
    suffix = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if content_type == "application/pdf" or suffix == ".pdf":
        return _extract_pdf_text(data)
    if content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document" or suffix == ".docx":
        document = Document(BytesIO(data))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        table_rows = [" | ".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows]
        return "\n".join([*paragraphs, *table_rows])
    if content_type == "text/plain" or suffix == ".txt":
        return data.decode("utf-8", errors="replace")
    raise HTTPException(status_code=415, detail="Upload a PDF, DOCX, or TXT resume")


@router.post("/import")
async def import_resume(file: UploadFile = File(...), current_user: dict = Depends(get_current_user)):
    require_feature(current_user, "resume_builder")
    filename = (file.filename or "resume").strip()[:240]
    data = await file.read(MAX_RESUME_BYTES + 1)
    if len(data) > MAX_RESUME_BYTES:
        raise HTTPException(status_code=413, detail="Resume file must be 5 MB or smaller")
    if not data:
        raise HTTPException(status_code=400, detail="Resume file is empty")
    try:
        text = _extract_uploaded_text(filename, file.content_type or "", data).strip()
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=400, detail="We could not read that resume. Try exporting it as a fresh PDF or DOCX.") from exc
    if len(text) < 20:
        raise HTTPException(
            status_code=400,
            detail="We could not find enough readable text in that resume. If it is a scanned/image-only PDF, export a text-searchable PDF or DOCX and try again.",
        )
    parsed = parse_existing_resume_text(text)
    return {
        "filename": filename,
        "text": parsed["text"],
        "raw_text": parsed.get("raw_text", text),
        "summary": parsed["summary"],
        "bullets": parsed["bullets"],
        "skills": parsed["skills"],
        "sections_found": parsed["sections_found"],
        "sections": parsed.get("sections", {}),
        "header_lines": parsed.get("header_lines", []),
        "contact": parsed.get("contact", {}),
        "experience": parsed.get("experience", []),
        "parse_warnings": parsed.get("parse_warnings", []),
        "source_signals": parsed.get("source_signals", {}),
        "line_count": parsed["line_count"],
    }


@router.post("/build")
def build_resume(payload: ResumeBuildRequest, current_user: dict = Depends(get_current_user)):
    require_feature(current_user, "resume_builder")
    receipts = _owned_receipts(str(current_user["_id"]), payload.selected_receipt_ids)
    result = analyze_resume(
        target_role=payload.target_role,
        job_description=payload.job_description,
        receipts=receipts,
        existing_resume_text=payload.existing_resume_text,
    )
    result["source_receipt_count"] = len(receipts)
    return result


@router.post("/resumes", status_code=status.HTTP_201_CREATED)
def save_resume(payload: ResumeSaveRequest, current_user: dict = Depends(get_current_user)):
    require_feature(current_user, "resume_builder")
    user_id = str(current_user["_id"])
    canonical_bullets = _canonical_saved_bullets(payload)
    if len(canonical_bullets) > 100:
        raise HTTPException(status_code=400, detail="A saved resume can contain up to 100 experience bullets")
    _validate_saved_bullet_sources(user_id, canonical_bullets)
    now = datetime.now(timezone.utc)
    bullets = [bullet.model_dump() for bullet in payload.bullets]
    experience = [role.model_dump() for role in payload.experience]
    supporting_sections = payload.supporting_sections.model_dump()
    document = {
        "user_id": user_id,
        "title": payload.title.strip(),
        "target_role": payload.target_role.strip(),
        "job_description": payload.job_description,
        "summary": payload.summary,
        "bullets": bullets,
        "skills": [skill.strip()[:120] for skill in payload.skills if skill.strip()],
        "contact": payload.contact.model_dump(),
        "experience": experience,
        "supporting_sections": supporting_sections,
        "readiness": _server_readiness(canonical_bullets),
        "schema_version": 4,
        "created_at": now,
        "updated_at": now,
    }
    result = resume_documents_collection.insert_one(document)
    return {"id": str(result.inserted_id), **{k: v for k, v in document.items() if k != "user_id"}}


@router.get("/resumes")
def list_resumes(current_user: dict = Depends(get_current_user)):
    require_feature(current_user, "resume_builder")
    cursor = resume_documents_collection.find({"user_id": str(current_user["_id"])}).sort("updated_at", -1).limit(50)
    return {
        "resumes": [
            {
                "id": str(item["_id"]),
                "title": item.get("title", "Untitled resume"),
                "target_role": item.get("target_role", ""),
                "job_description": item.get("job_description", ""),
                "summary": item.get("summary", ""),
                "bullets": item.get("bullets", []),
                "skills": item.get("skills", []),
                "contact": item.get("contact", {}),
                "experience": item.get("experience", []),
                "supporting_sections": item.get("supporting_sections", {}),
                "readiness": item.get("readiness", {}),
                "schema_version": item.get("schema_version", 3),
                "created_at": item.get("created_at"),
                "updated_at": item.get("updated_at"),
            }
            for item in cursor
        ]
    }


@router.delete("/resumes/{resume_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_resume(resume_id: str, current_user: dict = Depends(get_current_user)):
    require_feature(current_user, "resume_builder")
    if not ObjectId.is_valid(resume_id):
        raise HTTPException(status_code=400, detail="Invalid resume ID")
    result = resume_documents_collection.delete_one({"_id": ObjectId(resume_id), "user_id": str(current_user["_id"])})
    if not result.deleted_count:
        raise HTTPException(status_code=404, detail="Resume not found")
    return None
