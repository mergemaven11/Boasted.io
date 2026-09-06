"""Reviewer worksheets appended to Boasted career packet PDF and DOCX exports."""
from __future__ import annotations

from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pypdf import PdfReader, PdfWriter
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.packet_document_exports import PALETTES, build_career_packet_docx, build_career_packet_pdf


REVIEWER_GUIDES: dict[str, list[tuple[str, str]]] = {
    "performance-review": [
        ("Results and impact", "Are the outcomes clear, accurate, and meaningful for the review period?"),
        ("Ownership and contribution", "Is the person's specific role in the work clear?"),
        ("Collaboration and recognition", "Does the packet show how the work was carried with or recognized by others?"),
        ("Skills and growth", "Does the evidence show capability growth or repeated application?"),
        ("Evidence quality", "Are claims supported well enough to discuss confidently?"),
        ("Review narrative", "Does the packet tell a fair, complete story of the period?"),
    ],
    "promotion": [
        ("Expanded scope", "Does the evidence show broader responsibility, complexity, or ownership?"),
        ("Sustained results", "Are the strongest outcomes repeated or durable rather than one-off?"),
        ("Leadership and influence", "Does the record show influence, mentorship, coordination, or initiative where relevant?"),
        ("Next-level capabilities", "Do the documented skills match the expectations of the target level?"),
        ("Evidence quality", "Are the progression claims supported by concrete proof?"),
        ("Case clarity", "Is the progression case easy to follow without overstating readiness?"),
    ],
    "interview": [
        ("Role relevance", "Are the selected stories relevant to the role or organization?"),
        ("Story clarity", "Can the situation, contribution, result, and learning be understood quickly?"),
        ("Ownership", "Is the candidate's personal contribution distinct from the team's work?"),
        ("Outcome strength", "Are results specific and supported where possible?"),
        ("Evidence credibility", "Could the candidate defend the claims if asked for detail?"),
        ("Preparation gaps", "What missing context or follow-up questions should be prepared?"),
    ],
    "certification": [
        ("Competency alignment", "Does the evidence map clearly to the credential or licensure expectations?"),
        ("Experience relevance", "Is the documented experience directly relevant to the review?"),
        ("Evidence sufficiency", "Is there enough support for the competencies being claimed?"),
        ("Recency and continuity", "Is the work recent or sustained enough for the stated requirement?"),
        ("Accuracy and compliance", "Are claims precise and free of unsupported or confidential details?"),
        ("Requirement gaps", "What requirements still need stronger proof or clarification?"),
    ],
    "program-application": [
        ("Program alignment", "Does the packet connect the applicant's evidence to the program's focus?"),
        ("Growth and learning", "Does the record show learning, persistence, or development?"),
        ("Initiative and impact", "Are there concrete examples of ownership or meaningful contribution?"),
        ("Evidence strength", "Are the strongest statements supported by documented proof?"),
        ("Narrative clarity", "Does the material tell a coherent application story?"),
        ("Missing context", "What should the applicant clarify before submitting?"),
    ],
    "scholarship": [
        ("Selection alignment", "Does the evidence connect to the scholarship or award criteria?"),
        ("Leadership or service", "Where relevant, does the packet show initiative, service, or leadership?"),
        ("Documented impact", "Are outcomes specific enough to support the application?"),
        ("Recognition and evidence", "Are meaningful claims backed by evidence or recognition where available?"),
        ("Story clarity", "Does the material support a compelling but factual narrative?"),
        ("Missing context", "What needs clarification before an essay, nomination, or interview?"),
    ],
    "portfolio": [
        ("Audience relevance", "Are the selected projects appropriate for the intended audience?"),
        ("Project impact", "Do the examples explain what changed because of the work?"),
        ("Craft and quality", "Does the work demonstrate the quality expected for the field?"),
        ("Skills depth and breadth", "Does the portfolio show the right balance of capability?"),
        ("Evidence credibility", "Are outcomes and claims supported by available proof?"),
        ("Presentation clarity", "Can a reviewer quickly understand the strongest work?"),
    ],
    "career-transition": [
        ("Transferable skills", "Are the strongest transferable capabilities obvious?"),
        ("Target relevance", "Does the packet connect prior work to the target field or role?"),
        ("Evidence bridge", "Are the transition claims grounded in documented examples?"),
        ("Learning and growth", "Does the record show preparation for the new direction?"),
        ("Narrative credibility", "Is the transition story persuasive without claiming undocumented experience?"),
        ("Gaps to address", "What skills, context, or proof should be strengthened next?"),
    ],
}

FEEDBACK_SECTIONS = [
    "Strengths / what stands out",
    "Corrections or factual changes",
    "Questions / clarification needed",
    "Recommended next steps",
]


def reviewer_guide_for_packet(packet_type: str) -> dict[str, Any]:
    criteria = REVIEWER_GUIDES.get(packet_type, REVIEWER_GUIDES["performance-review"])
    return {
        "rating_scale": "1 = needs clarification; 3 = solid; 5 = especially strong; N/A = not applicable",
        "criteria": [{"label": label, "prompt": prompt} for label, prompt in criteria],
        "feedback_sections": list(FEEDBACK_SECTIONS),
        "source_record_notice": "Reviewer corrections are suggestions and do not automatically change Boasted source records.",
    }


def _palette(packet: dict[str, Any]) -> dict[str, str]:
    name = packet.get("render_config", {}).get("theme") or "modern-minimal"
    return PALETTES.get(name, PALETTES["modern-minimal"])


def _pdf_reviewer_page(packet: dict[str, Any]) -> bytes:
    palette = _palette(packet)
    primary = colors.HexColor(f"#{palette['primary']}")
    secondary = colors.HexColor(f"#{palette['secondary']}")
    soft = colors.HexColor(f"#{palette['soft']}")
    line = colors.HexColor(f"#{palette['accent']}")
    ink = colors.HexColor(f"#{palette['ink']}")
    muted = colors.HexColor(f"#{palette['muted']}")
    guide = packet.get("reviewer_guide") or reviewer_guide_for_packet(str(packet.get("kind") or "performance-review"))
    criteria = (guide.get("criteria") or [])[:6]
    reviewer = str((packet.get("branding") or {}).get("reviewer_name") or "").strip()

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        leftMargin=0.62 * inch,
        rightMargin=0.62 * inch,
        topMargin=0.55 * inch,
        bottomMargin=0.48 * inch,
        title="Boasted Reviewer Worksheet",
        author="Boasted",
    )
    base = getSampleStyleSheet()
    eyebrow = ParagraphStyle("ReviewerEyebrow", parent=base["Normal"], fontName="Helvetica-Bold", fontSize=7.5, leading=9, textColor=secondary, spaceAfter=5)
    title = ParagraphStyle("ReviewerTitle", parent=base["Heading1"], fontName="Helvetica-Bold", fontSize=22, leading=25, textColor=primary, spaceAfter=7)
    body = ParagraphStyle("ReviewerBody", parent=base["BodyText"], fontName="Helvetica", fontSize=7.8, leading=10.4, textColor=muted, spaceAfter=5)
    small = ParagraphStyle("ReviewerSmall", parent=base["BodyText"], fontName="Helvetica", fontSize=6.4, leading=8.2, textColor=ink)
    small_bold = ParagraphStyle("ReviewerSmallBold", parent=small, fontName="Helvetica-Bold", textColor=primary)

    story = [
        Paragraph("10 · REVIEWER WORKSHEET", eyebrow),
        Paragraph("Grade the evidence. Leave the record better.", title),
        Paragraph("Use this worksheet to assess the packet, flag factual corrections, and leave actionable feedback. Ratings are completed by the reviewer—not generated by Boasted.", body),
    ]

    meta = Table([
        [Paragraph("Reviewer", small_bold), Paragraph(reviewer or "________________________", small), Paragraph("Date", small_bold), Paragraph("____________", small)],
        [Paragraph("Role / relationship", small_bold), Paragraph("________________________", small), Paragraph("Overall assessment", small_bold), Paragraph("____ / 5 or N/A", small)],
    ], colWidths=[1.15 * inch, 2.35 * inch, 1.25 * inch, 1.55 * inch])
    meta.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, -1), soft),
        ("BACKGROUND", (2, 0), (2, -1), soft),
        ("GRID", (0, 0), (-1, -1), 0.45, line),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.extend([meta, Spacer(1, 8)])

    rows = [[Paragraph("Section to review", small_bold), Paragraph("Grade", small_bold), Paragraph("Reviewer note", small_bold)]]
    for item in criteria:
        label = str(item.get("label") or "Review area")
        prompt = str(item.get("prompt") or "")
        rows.append([
            Paragraph(f"<b>{label}</b><br/><font color='#{palette['muted']}'>{prompt}</font>", small),
            Paragraph("____ / 5", small_bold),
            Paragraph("________________________________", small),
        ])
    rubric = Table(rows, colWidths=[3.55 * inch, 0.8 * inch, 2.0 * inch], repeatRows=1)
    rubric.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), primary),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.4, line),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 1), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 1), (-1, -1), 5),
    ]))
    story.extend([rubric, Spacer(1, 8)])

    feedback_cells = []
    for label in FEEDBACK_SECTIONS:
        feedback_cells.append([
            Paragraph(label, small_bold),
            Paragraph("______________________________________________<br/>______________________________________________", small),
        ])
    feedback = Table([
        [feedback_cells[0], feedback_cells[1]],
        [feedback_cells[2], feedback_cells[3]],
    ], colWidths=[3.18 * inch, 3.18 * inch])
    feedback.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.white),
        ("BOX", (0, 0), (-1, -1), 0.45, line),
        ("INNERGRID", (0, 0), (-1, -1), 0.35, line),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 7),
        ("RIGHTPADDING", (0, 0), (-1, -1), 7),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    story.extend([
        feedback,
        Spacer(1, 7),
        Paragraph("<b>Corrections are suggestions, not automatic edits.</b> Reviewer notes do not overwrite accomplishments, Impact Receipts, evidence, or recognition in Boasted. The packet owner decides what source records should be updated.", body),
    ])
    doc.build(story)
    return buffer.getvalue()


def build_reviewable_career_packet_pdf(packet: dict[str, Any]) -> bytes:
    base_bytes = build_career_packet_pdf(packet)
    reviewer_bytes = _pdf_reviewer_page(packet)
    writer = PdfWriter()
    for source in (base_bytes, reviewer_bytes):
        reader = PdfReader(BytesIO(source))
        for page in reader.pages:
            writer.add_page(page)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def _set_docx_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _format_docx_run(run, *, size: float, bold: bool = False, color: str = "17202A") -> None:
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def _add_feedback_box(document: Document, label: str, palette: dict[str, str]) -> None:
    table = document.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    _set_docx_cell_shading(cell, palette["soft"])
    paragraph = cell.paragraphs[0]
    label_run = paragraph.add_run(label)
    _format_docx_run(label_run, size=9, bold=True, color=palette["primary"])
    for _ in range(3):
        line = cell.add_paragraph("____________________________________________________________")
        for run in line.runs:
            _format_docx_run(run, size=8, color=palette["muted"])


def build_reviewable_career_packet_docx(packet: dict[str, Any]) -> bytes:
    base_bytes = build_career_packet_docx(packet)
    document = Document(BytesIO(base_bytes))
    palette = _palette(packet)
    guide = packet.get("reviewer_guide") or reviewer_guide_for_packet(str(packet.get("kind") or "performance-review"))
    criteria = (guide.get("criteria") or [])[:6]
    reviewer = str((packet.get("branding") or {}).get("reviewer_name") or "").strip()

    document.add_page_break()
    eyebrow = document.add_paragraph()
    eyebrow_run = eyebrow.add_run("REVIEWER WORKSHEET")
    _format_docx_run(eyebrow_run, size=8, bold=True, color=palette["secondary"])
    heading = document.add_paragraph()
    heading_run = heading.add_run("Grade the evidence. Leave the record better.")
    _format_docx_run(heading_run, size=21, bold=True, color=palette["primary"])
    intro = document.add_paragraph("Use this worksheet to assess the packet, flag factual corrections, and leave actionable feedback. Ratings are completed by the reviewer—not generated by Boasted.")
    for run in intro.runs:
        _format_docx_run(run, size=9, color=palette["muted"])

    meta = document.add_table(rows=2, cols=4)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_values = [
        ("Reviewer", reviewer or "________________________", "Date", "____________"),
        ("Role / relationship", "________________________", "Overall assessment", "____ / 5 or N/A"),
    ]
    for row_index, row_values in enumerate(meta_values):
        for col_index, value in enumerate(row_values):
            cell = meta.cell(row_index, col_index)
            if col_index in (0, 2):
                _set_docx_cell_shading(cell, palette["soft"])
            run = cell.paragraphs[0].add_run(value)
            _format_docx_run(run, size=8.5, bold=col_index in (0, 2), color=palette["primary"] if col_index in (0, 2) else palette["ink"])

    document.add_paragraph()
    rubric = document.add_table(rows=1, cols=3)
    rubric.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Section to review", "Grade", "Reviewer note"]
    for index, label in enumerate(headers):
        _set_docx_cell_shading(rubric.cell(0, index), palette["primary"])
        run = rubric.cell(0, index).paragraphs[0].add_run(label)
        _format_docx_run(run, size=8, bold=True, color="FFFFFF")
    for item in criteria:
        cells = rubric.add_row().cells
        label = str(item.get("label") or "Review area")
        prompt = str(item.get("prompt") or "")
        label_run = cells[0].paragraphs[0].add_run(label)
        _format_docx_run(label_run, size=8.5, bold=True, color=palette["primary"])
        prompt_run = cells[0].add_paragraph(prompt).add_run()
        _format_docx_run(prompt_run, size=7.5, color=palette["muted"])
        grade_run = cells[1].paragraphs[0].add_run("____ / 5")
        _format_docx_run(grade_run, size=8.5, bold=True, color=palette["primary"])
        note_run = cells[2].paragraphs[0].add_run("____________________________")
        _format_docx_run(note_run, size=8, color=palette["muted"])

    document.add_paragraph()
    for label in FEEDBACK_SECTIONS:
        _add_feedback_box(document, label, palette)
        document.add_paragraph()

    safety = document.add_paragraph()
    lead = safety.add_run("Corrections are suggestions, not automatic edits. ")
    _format_docx_run(lead, size=8.5, bold=True, color=palette["primary"])
    body = safety.add_run("Reviewer notes do not overwrite accomplishments, Impact Receipts, evidence, or recognition in Boasted. The packet owner decides what source records should be updated.")
    _format_docx_run(body, size=8.5, color=palette["muted"])

    section = document.sections[-1]
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.62)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    output = BytesIO()
    document.save(output)
    return output.getvalue()
