"""Themed PDF and DOCX exports for BragStack career packets."""
from __future__ import annotations

import io
import re
from typing import Any
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    KeepTogether,
    PageTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)


PALETTES = {
    "classic-dossier": {
        "primary": "173F43",
        "secondary": "2B6F73",
        "accent": "9BDDE0",
        "soft": "EFF7F6",
        "ink": "17202A",
        "muted": "66727A",
    },
    "modern-minimal": {
        "primary": "202A55",
        "secondary": "5E57C8",
        "accent": "8ED8F8",
        "soft": "F3F3FF",
        "ink": "17202A",
        "muted": "667085",
    },
    "executive-report": {
        "primary": "202A36",
        "secondary": "315D8C",
        "accent": "C9A44B",
        "soft": "F5F6F8",
        "ink": "1B2631",
        "muted": "66727A",
    },
}

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _clean(value: Any, fallback: str = "") -> str:
    text = str(value or "").strip()
    return text or fallback


def _hex_color(value: str) -> colors.Color:
    return colors.HexColor(f"#{value}")


def _theme(packet: dict[str, Any]) -> dict[str, str]:
    name = packet.get("render_config", {}).get("theme") or "modern-minimal"
    return PALETTES.get(name, PALETTES["modern-minimal"])


def _slug(value: str, fallback: str = "packet") -> str:
    text = re.sub(r"[^a-z0-9]+", "-", (value or "").lower()).strip("-")
    return text[:80] or fallback


def make_career_packet_filename(packet: dict[str, Any], extension: str) -> str:
    subject = _clean(packet.get("subject", {}).get("name"), "bragstack-member")
    kind = _clean(packet.get("kind"), "career-packet")
    ext = "docx" if extension.lower() == "docx" else "pdf"
    return f"bragstack-{_slug(subject)}-{_slug(kind)}.{ext}"


def _focus_fields(packet: dict[str, Any]) -> list[dict[str, str]]:
    fields = packet.get("focus_fields") or []
    result = []
    for item in fields:
        label = _clean(item.get("label")) if isinstance(item, dict) else ""
        value = _clean(item.get("value")) if isinstance(item, dict) else ""
        if label and value:
            result.append({"label": label, "value": value})
    return result[:10]


def _top_skills(packet: dict[str, Any]) -> list[tuple[str, int]]:
    items = packet.get("impact_analytics", {}).get("top_skills") or {}
    if isinstance(items, dict):
        return [(str(name), int(count or 0)) for name, count in list(items.items())[:10]]
    return []


def _paragraph_text(value: Any) -> str:
    return escape(_clean(value))


def _pdf_styles(palette: dict[str, str]):
    base = getSampleStyleSheet()
    return {
        "eyebrow": ParagraphStyle(
            "PacketEyebrow",
            parent=base["Normal"],
            fontName="Helvetica-Bold",
            fontSize=8,
            leading=10,
            textColor=_hex_color(palette["secondary"]),
            spaceAfter=5,
        ),
        "title": ParagraphStyle(
            "PacketTitle",
            parent=base["Title"],
            fontName="Helvetica-Bold",
            fontSize=28,
            leading=32,
            textColor=_hex_color(palette["primary"]),
            spaceAfter=8,
        ),
        "subtitle": ParagraphStyle(
            "PacketSubtitle",
            parent=base["Normal"],
            fontName="Helvetica",
            fontSize=12,
            leading=17,
            textColor=_hex_color(palette["muted"]),
            spaceAfter=9,
        ),
        "h2": ParagraphStyle(
            "PacketH2",
            parent=base["Heading2"],
            fontName="Helvetica-Bold",
            fontSize=17,
            leading=21,
            textColor=_hex_color(palette["primary"]),
            spaceBefore=5,
            spaceAfter=9,
        ),
        "h3": ParagraphStyle(
            "PacketH3",
            parent=base["Heading3"],
            fontName="Helvetica-Bold",
            fontSize=11,
            leading=14,
            textColor=_hex_color(palette["ink"]),
            spaceAfter=4,
        ),
        "body": ParagraphStyle(
            "PacketBody",
            parent=base["BodyText"],
            fontName="Helvetica",
            fontSize=9.5,
            leading=14,
            textColor=_hex_color(palette["ink"]),
            spaceAfter=7,
        ),
        "small": ParagraphStyle(
            "PacketSmall",
            parent=base["BodyText"],
            fontName="Helvetica",
            fontSize=8,
            leading=11,
            textColor=_hex_color(palette["muted"]),
            spaceAfter=4,
        ),
        "stat": ParagraphStyle(
            "PacketStat",
            parent=base["Normal"],
            fontName="Helvetica-Bold",
            fontSize=19,
            leading=22,
            alignment=TA_CENTER,
            textColor=_hex_color(palette["primary"]),
        ),
        "statLabel": ParagraphStyle(
            "PacketStatLabel",
            parent=base["Normal"],
            fontName="Helvetica",
            fontSize=7.5,
            leading=9,
            alignment=TA_CENTER,
            textColor=_hex_color(palette["muted"]),
        ),
    }


def _pdf_page(canvas, doc, palette: dict[str, str]):
    canvas.saveState()
    width, height = letter
    canvas.setFillColor(_hex_color(palette["primary"]))
    canvas.rect(0, height - 0.25 * inch, width, 0.25 * inch, stroke=0, fill=1)
    canvas.setFillColor(_hex_color(palette["secondary"]))
    canvas.rect(0, 0, width, 0.08 * inch, stroke=0, fill=1)
    canvas.setFont("Helvetica", 7)
    canvas.setFillColor(_hex_color(palette["muted"]))
    canvas.drawString(0.65 * inch, 0.27 * inch, "BragStack - Career Evidence System")
    canvas.drawRightString(width - 0.65 * inch, 0.27 * inch, f"Page {doc.page}")
    canvas.restoreState()


def _pdf_section_header(story: list, title: str, styles: dict[str, ParagraphStyle]):
    story.append(Spacer(1, 6))
    story.append(Paragraph(_paragraph_text(title), styles["h2"]))


def _pdf_item_card(item: dict[str, Any], styles, palette):
    title = _clean(item.get("title") or item.get("accomplishment"), "Documented accomplishment")
    category = _clean(item.get("category"))
    result = _clean(item.get("result"))
    skills = item.get("skills") or []
    pieces = [Paragraph(_paragraph_text(title), styles["h3"])]
    if category:
        pieces.append(Paragraph(_paragraph_text(category), styles["small"]))
    if result:
        pieces.append(Paragraph(_paragraph_text(result), styles["body"]))
    if skills:
        pieces.append(Paragraph(f"<b>Skills:</b> {_paragraph_text(', '.join(str(skill) for skill in skills[:8]))}", styles["small"]))
    table = Table([[pieces]], colWidths=[6.65 * inch])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), _hex_color(palette["soft"])),
        ("BOX", (0, 0), (-1, -1), 0.5, _hex_color(palette["accent"])),
        ("LEFTPADDING", (0, 0), (-1, -1), 10),
        ("RIGHTPADDING", (0, 0), (-1, -1), 10),
        ("TOPPADDING", (0, 0), (-1, -1), 9),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
    ]))
    return table


def build_career_packet_pdf(packet: dict[str, Any]) -> bytes:
    palette = _theme(packet)
    styles = _pdf_styles(palette)
    buffer = io.BytesIO()
    doc = BaseDocTemplate(
        buffer,
        pagesize=letter,
        leftMargin=0.65 * inch,
        rightMargin=0.65 * inch,
        topMargin=0.58 * inch,
        bottomMargin=0.55 * inch,
        title=_clean(packet.get("title"), "BragStack Career Packet"),
        author="BragStack",
    )
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="normal")
    doc.addPageTemplates([PageTemplate(id="packet", frames=frame, onPage=lambda c, d: _pdf_page(c, d, palette))])

    title = _clean(packet.get("title"), "Career Packet")
    subject = packet.get("subject") or {}
    scorecard = packet.get("scorecard") or {}
    branding = packet.get("branding") or {}
    context = packet.get("context") or {}
    focus = _focus_fields(packet)
    story = []

    story.append(Spacer(1, 14))
    story.append(Paragraph("BRAGSTACK PROFESSIONAL PACKET", styles["eyebrow"]))
    story.append(Paragraph(_paragraph_text(title), styles["title"]))
    story.append(Paragraph(_paragraph_text(_clean(subject.get("name"), "BragStack Member")), styles["h2"]))
    role_line = " - ".join(value for value in [_clean(subject.get("role")), _clean(context.get("organization"))] if value)
    if role_line:
        story.append(Paragraph(_paragraph_text(role_line), styles["subtitle"]))
    if branding.get("brand_name"):
        story.append(Paragraph(f"Prepared under <b>{_paragraph_text(branding.get('brand_name'))}</b>", styles["small"]))

    if focus:
        focus_rows = [[Paragraph(_paragraph_text(item["label"]), styles["small"]), Paragraph(_paragraph_text(item["value"]), styles["body"])] for item in focus]
        focus_table = Table(focus_rows, colWidths=[1.55 * inch, 5.1 * inch])
        focus_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (0, -1), _hex_color(palette["soft"])),
            ("BOX", (0, 0), (-1, -1), 0.5, _hex_color(palette["accent"])),
            ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#E1E5EC")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 8),
            ("RIGHTPADDING", (0, 0), (-1, -1), 8),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ]))
        story.append(Spacer(1, 8))
        story.append(focus_table)

    stats = [
        (scorecard.get("accomplishments", 0), "Accomplishments"),
        (scorecard.get("impact_receipts", 0), "Impact Receipts"),
        (scorecard.get("evidence_items", 0), "Evidence Items"),
        (scorecard.get("skills_demonstrated", 0), "Skills"),
    ]
    stats_table = Table([[Paragraph(str(value), styles["stat"]) for value, _ in stats], [Paragraph(label, styles["statLabel"]) for _, label in stats]], colWidths=[1.65 * inch] * 4)
    stats_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), _hex_color(palette["soft"])),
        ("BOX", (0, 0), (-1, -1), 0.6, _hex_color(palette["accent"])),
        ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#E2E6EC")),
        ("TOPPADDING", (0, 0), (-1, 0), 9),
        ("BOTTOMPADDING", (0, 1), (-1, 1), 9),
    ]))
    story.append(Spacer(1, 13))
    story.append(stats_table)

    quality = _clean(packet.get("quality_message"))
    if quality:
        quality_table = Table([[Paragraph(f"<b>Build a stronger packet:</b> {_paragraph_text(quality)}", styles["body"])]], colWidths=[6.6 * inch])
        quality_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), _hex_color(palette["soft"])),
            ("BOX", (0, 0), (-1, -1), 0.7, _hex_color(palette["secondary"])),
            ("LEFTPADDING", (0, 0), (-1, -1), 10),
            ("RIGHTPADDING", (0, 0), (-1, -1), 10),
            ("TOPPADDING", (0, 0), (-1, -1), 8),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
        ]))
        story.append(Spacer(1, 11))
        story.append(quality_table)

    _pdf_section_header(story, "Purpose and summary", styles)
    usage = _clean(packet.get("usage_example"))
    if usage:
        story.append(Paragraph(f"<b>How this packet can be used:</b> {_paragraph_text(usage)}", styles["body"]))
    story.append(Paragraph(_paragraph_text(_clean(packet.get("review_summary"), "This packet organizes the career proof currently saved in BragStack.")), styles["body"]))

    signature = packet.get("signature_accomplishments") or []
    _pdf_section_header(story, "Featured accomplishments", styles)
    if signature:
        for item in signature[:8]:
            story.append(KeepTogether([_pdf_item_card(item, styles, palette), Spacer(1, 7)]))
    else:
        story.append(Paragraph("No accomplishments are available yet. Add accomplishments and Impact Receipts to populate this section.", styles["body"]))

    measurable = packet.get("measurable_results") or []
    if measurable:
        _pdf_section_header(story, "Measurable results", styles)
        for item in measurable[:8]:
            result = _clean(item.get("result"))
            story.append(Paragraph(f"<b>{_paragraph_text(item.get('title') or 'Result')}</b> - {_paragraph_text(result)}", styles["body"]))

    skills = _top_skills(packet)
    if skills:
        _pdf_section_header(story, "Demonstrated skills", styles)
        skill_rows = [[Paragraph(_paragraph_text(name), styles["body"]), Paragraph(str(count), styles["body"])] for name, count in skills]
        skill_table = Table(skill_rows, colWidths=[5.7 * inch, 0.9 * inch])
        skill_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.white),
            ("ROWBACKGROUNDS", (0, 0), (-1, -1), [colors.white, _hex_color(palette["soft"])]),
            ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#E1E5EC")),
            ("LEFTPADDING", (0, 0), (-1, -1), 7),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        story.append(skill_table)

    receipts = packet.get("receipt_records") or []
    if receipts:
        _pdf_section_header(story, "Impact Receipts and recognition", styles)
        for receipt in receipts[:10]:
            recognition = receipt.get("recognition") or []
            recognition_text = ", ".join(_clean(item.get("label")) for item in recognition if _clean(item.get("label")))
            text = _clean(receipt.get("result") or receipt.get("contribution"))
            extra = f" Verified Recognition: {recognition_text}." if recognition_text else ""
            story.append(Paragraph(f"<b>{_paragraph_text(receipt.get('accomplishment') or 'Impact Receipt')}</b> - {_paragraph_text(text + extra)}", styles["body"]))

    evidence = packet.get("evidence_index") or []
    if evidence:
        _pdf_section_header(story, "Evidence index", styles)
        for item in evidence[:20]:
            label = _clean(item.get("title"), "Evidence item")
            reference = _clean(item.get("reference"))
            line = f"{label}{' - ' + reference if reference else ''}"
            story.append(Paragraph(_paragraph_text(line), styles["small"]))

    story.append(Spacer(1, 10))
    story.append(Paragraph("Generated from user-saved BragStack data. Review the document before submitting it to an employer, school, licensing body, scholarship committee, client, or other third party.", styles["small"]))
    doc.build(story)
    return buffer.getvalue()


def _set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _docx_run(run, *, size=None, bold=None, color=None, font="Aptos"):
    run.font.name = font
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _docx_add_heading(document: Document, text: str, palette: dict[str, str], level: int = 1):
    paragraph = document.add_paragraph()
    paragraph.style = document.styles["Heading 1" if level == 1 else "Heading 2"]
    run = paragraph.add_run(text)
    _docx_run(run, size=17 if level == 1 else 12, bold=True, color=palette["primary"])
    return paragraph


def _docx_add_body(document: Document, text: str, palette: dict[str, str], bold_prefix: str | None = None):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    if bold_prefix:
        prefix = paragraph.add_run(bold_prefix)
        _docx_run(prefix, size=10, bold=True, color=palette["primary"])
    run = paragraph.add_run(text)
    _docx_run(run, size=10, color=palette["ink"])
    return paragraph


def build_career_packet_docx(packet: dict[str, Any]) -> bytes:
    palette = _theme(packet)
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.68)
    section.right_margin = Inches(0.68)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header.add_run("BRAGSTACK  |  CAREER EVIDENCE SYSTEM")
    _docx_run(header_run, size=8, bold=True, color=palette["secondary"])
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("Generated from user-saved BragStack proof - review before external submission")
    _docx_run(footer_run, size=7.5, color=palette["muted"])

    title = _clean(packet.get("title"), "Career Packet")
    subject = packet.get("subject") or {}
    context = packet.get("context") or {}
    branding = packet.get("branding") or {}
    scorecard = packet.get("scorecard") or {}

    eyebrow = document.add_paragraph()
    eyebrow_run = eyebrow.add_run("BRAGSTACK PROFESSIONAL PACKET")
    _docx_run(eyebrow_run, size=8, bold=True, color=palette["secondary"])
    title_paragraph = document.add_paragraph()
    title_run = title_paragraph.add_run(title)
    _docx_run(title_run, size=25, bold=True, color=palette["primary"])
    title_paragraph.paragraph_format.space_after = Pt(3)
    subject_paragraph = document.add_paragraph()
    subject_run = subject_paragraph.add_run(_clean(subject.get("name"), "BragStack Member"))
    _docx_run(subject_run, size=17, bold=True, color=palette["ink"])
    role_line = " - ".join(value for value in [_clean(subject.get("role")), _clean(context.get("organization"))] if value)
    if role_line:
        role_paragraph = document.add_paragraph()
        role_run = role_paragraph.add_run(role_line)
        _docx_run(role_run, size=11, color=palette["muted"])
    if branding.get("brand_name"):
        _docx_add_body(document, _clean(branding.get("brand_name")), palette, "Cover label: ")

    focus = _focus_fields(packet)
    if focus:
        table = document.add_table(rows=0, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        for item in focus:
            cells = table.add_row().cells
            cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_shading(cells[0], palette["soft"])
            for cell in cells:
                _set_cell_margins(cell)
            left = cells[0].paragraphs[0].add_run(item["label"])
            _docx_run(left, size=8.5, bold=True, color=palette["secondary"])
            right = cells[1].paragraphs[0].add_run(item["value"])
            _docx_run(right, size=9.5, color=palette["ink"])
        document.add_paragraph()

    stat_table = document.add_table(rows=2, cols=4)
    stat_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    stat_table.autofit = True
    stats = [
        (scorecard.get("accomplishments", 0), "Accomplishments"),
        (scorecard.get("impact_receipts", 0), "Impact Receipts"),
        (scorecard.get("evidence_items", 0), "Evidence Items"),
        (scorecard.get("skills_demonstrated", 0), "Skills"),
    ]
    for index, (value, label) in enumerate(stats):
        for row in (0, 1):
            _set_cell_shading(stat_table.cell(row, index), palette["soft"])
            _set_cell_margins(stat_table.cell(row, index), top=70, bottom=70)
        value_paragraph = stat_table.cell(0, index).paragraphs[0]
        value_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        value_run = value_paragraph.add_run(str(value))
        _docx_run(value_run, size=17, bold=True, color=palette["primary"])
        label_paragraph = stat_table.cell(1, index).paragraphs[0]
        label_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        label_run = label_paragraph.add_run(label)
        _docx_run(label_run, size=7.5, color=palette["muted"])

    quality = _clean(packet.get("quality_message"))
    if quality:
        document.add_paragraph()
        quality_table = document.add_table(rows=1, cols=1)
        cell = quality_table.cell(0, 0)
        _set_cell_shading(cell, palette["soft"])
        _set_cell_margins(cell, top=130, bottom=130, start=150, end=150)
        paragraph = cell.paragraphs[0]
        strong = paragraph.add_run("Build a stronger packet: ")
        _docx_run(strong, size=9.5, bold=True, color=palette["primary"])
        body = paragraph.add_run(quality)
        _docx_run(body, size=9.5, color=palette["ink"])

    _docx_add_heading(document, "Purpose and summary", palette)
    usage = _clean(packet.get("usage_example"))
    if usage:
        _docx_add_body(document, usage, palette, "How this packet can be used: ")
    _docx_add_body(document, _clean(packet.get("review_summary"), "This packet organizes the career proof currently saved in BragStack."), palette)

    _docx_add_heading(document, "Featured accomplishments", palette)
    signature = packet.get("signature_accomplishments") or []
    if signature:
        for item in signature[:8]:
            paragraph = document.add_paragraph(style="List Bullet")
            title_run = paragraph.add_run(_clean(item.get("title") or item.get("accomplishment"), "Documented accomplishment"))
            _docx_run(title_run, size=10, bold=True, color=palette["primary"])
            result = _clean(item.get("result"))
            if result:
                result_run = paragraph.add_run(f" - {result}")
                _docx_run(result_run, size=10, color=palette["ink"])
            skills = item.get("skills") or []
            if skills:
                skill_paragraph = document.add_paragraph()
                skill_paragraph.paragraph_format.left_indent = Inches(0.24)
                skill_run = skill_paragraph.add_run("Skills: " + ", ".join(str(skill) for skill in skills[:8]))
                _docx_run(skill_run, size=8.5, color=palette["muted"])
    else:
        _docx_add_body(document, "No accomplishments are available yet. Add accomplishments and Impact Receipts to populate this section.", palette)

    measurable = packet.get("measurable_results") or []
    if measurable:
        _docx_add_heading(document, "Measurable results", palette)
        for item in measurable[:8]:
            _docx_add_body(document, _clean(item.get("result")), palette, f"{_clean(item.get('title'), 'Result')}: ")

    skills_data = _top_skills(packet)
    if skills_data:
        _docx_add_heading(document, "Demonstrated skills", palette)
        table = document.add_table(rows=1, cols=2)
        table.style = "Light Shading Accent 1"
        headers = table.rows[0].cells
        headers[0].text = "Skill"
        headers[1].text = "Documented uses"
        for name, count in skills_data:
            cells = table.add_row().cells
            cells[0].text = name
            cells[1].text = str(count)
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        _docx_run(run, size=9, color=palette["ink"])

    receipts = packet.get("receipt_records") or []
    if receipts:
        _docx_add_heading(document, "Impact Receipts and recognition", palette)
        for receipt in receipts[:10]:
            recognition = receipt.get("recognition") or []
            recognition_text = ", ".join(_clean(item.get("label")) for item in recognition if _clean(item.get("label")))
            text = _clean(receipt.get("result") or receipt.get("contribution"))
            if recognition_text:
                text += f" Verified Recognition: {recognition_text}."
            _docx_add_body(document, text, palette, f"{_clean(receipt.get('accomplishment'), 'Impact Receipt')}: ")

    evidence = packet.get("evidence_index") or []
    if evidence:
        _docx_add_heading(document, "Evidence index", palette)
        for item in evidence[:20]:
            line = _clean(item.get("title"), "Evidence item")
            reference = _clean(item.get("reference"))
            if reference:
                line += f" - {reference}"
            paragraph = document.add_paragraph(style="List Bullet")
            run = paragraph.add_run(line)
            _docx_run(run, size=9, color=palette["ink"])

    document.add_paragraph()
    notice = document.add_paragraph()
    notice_run = notice.add_run("Generated from user-saved BragStack data. Review the document before submitting it to an employer, school, licensing body, scholarship committee, client, or other third party.")
    _docx_run(notice_run, size=8, color=palette["muted"])

    core = document.core_properties
    core.title = title
    core.subject = "BragStack career packet"
    core.author = "BragStack"
    core.comments = "Generated from user-saved BragStack career evidence."

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
